import os
import io
from typing import List, Tuple


class DocumentProcessor:
    """
    Processes uploaded files into plain text chunks.
    Supports: PDF, DOCX, TXT, CSV, XLSX
    """

    def process(self, file) -> Tuple[List[str], str]:
        """
        Takes a Streamlit UploadedFile, returns (list_of_text_chunks, filename).
        """
        name = file.name
        ext = name.split(".")[-1].lower()

        if ext == "pdf":
            return self._process_pdf(file, name)
        elif ext == "docx":
            return self._process_docx(file, name)
        elif ext == "txt":
            return self._process_txt(file, name)
        elif ext == "csv":
            return self._process_csv(file, name)
        elif ext == "xlsx":
            return self._process_xlsx(file, name)
        else:
            raw = file.read().decode("utf-8", errors="ignore")
            return [raw], name

    def _process_pdf(self, file, name: str) -> Tuple[List[str], str]:
        try:
            import pdfplumber
            texts = []
            with pdfplumber.open(io.BytesIO(file.read())) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text and text.strip():
                        # Add page reference for citation
                        texts.append(f"[Source: {name}, Page {page.page_number}]\n{text.strip()}")
            return texts if texts else ["(No extractable text found in PDF)"], name
        except ImportError:
            # Fallback: PyPDF2
            try:
                import PyPDF2
                reader = PyPDF2.PdfReader(io.BytesIO(file.read()))
                texts = []
                for i, page in enumerate(reader.pages):
                    text = page.extract_text()
                    if text and text.strip():
                        texts.append(f"[Source: {name}, Page {i+1}]\n{text.strip()}")
                return texts if texts else ["(No text extracted)"], name
            except Exception as e:
                return [f"Error reading PDF: {str(e)}"], name

    def _process_docx(self, file, name: str) -> Tuple[List[str], str]:
        try:
            from docx import Document
            doc = Document(io.BytesIO(file.read()))
            texts = []
            current_section = []

            for para in doc.paragraphs:
                if para.style.name.startswith("Heading") and current_section:
                    section_text = "\n".join(current_section)
                    texts.append(f"[Source: {name}]\n{section_text}")
                    current_section = []
                if para.text.strip():
                    current_section.append(para.text.strip())

            if current_section:
                texts.append(f"[Source: {name}]\n" + "\n".join(current_section))

            # Also extract tables
            for table in doc.tables:
                rows = []
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    rows.append(row_text)
                if rows:
                    texts.append(f"[Source: {name}, Table]\n" + "\n".join(rows))

            return texts if texts else ["(Empty document)"], name
        except Exception as e:
            return [f"Error reading DOCX: {str(e)}"], name

    def _process_txt(self, file, name: str) -> Tuple[List[str], str]:
        try:
            content = file.read().decode("utf-8", errors="ignore")
            # Split into paragraphs
            paragraphs = [p.strip() for p in content.split("\n\n") if p.strip()]
            texts = [f"[Source: {name}]\n{p}" for p in paragraphs]
            return texts if texts else [f"[Source: {name}]\n{content}"], name
        except Exception as e:
            return [f"Error reading TXT: {str(e)}"], name

    def _process_csv(self, file, name: str) -> Tuple[List[str], str]:
        try:
            import pandas as pd
            df = pd.read_csv(io.BytesIO(file.read()))
            texts = []
            # Convert in batches of 20 rows for better chunking
            batch_size = 20
            for i in range(0, len(df), batch_size):
                batch = df.iloc[i:i+batch_size]
                text = f"[Source: {name}, Rows {i+1}-{i+len(batch)}]\n"
                text += batch.to_string(index=False)
                texts.append(text)
            # Also add column summary
            summary = f"[Source: {name}, Summary]\nColumns: {', '.join(df.columns)}\nTotal rows: {len(df)}"
            texts.insert(0, summary)
            return texts, name
        except Exception as e:
            return [f"Error reading CSV: {str(e)}"], name

    def _process_xlsx(self, file, name: str) -> Tuple[List[str], str]:
        try:
            import pandas as pd
            xl = pd.ExcelFile(io.BytesIO(file.read()))
            texts = []
            for sheet_name in xl.sheet_names:
                df = xl.parse(sheet_name)
                texts.append(f"[Source: {name}, Sheet: {sheet_name}]\n{df.to_string(index=False)}")
            return texts if texts else ["(Empty Excel file)"], name
        except Exception as e:
            return [f"Error reading XLSX: {str(e)}"], name
